"""CSV and XLSX exports."""

from __future__ import annotations

import csv
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.models.schemas import PaymentRow
from app.services.reconciliation import overdue_rows

OVERDUE_COLUMNS: tuple[str, ...] = (
    "payment_id",
    "provider",
    "payment_date",
    "amount",
    "currency",
    "expected_settlement_date",
    "days_overdue",
    "status",
)

ALL_COLUMNS: tuple[str, ...] = (
    "payment_id",
    "provider",
    "payment_date",
    "amount",
    "currency",
    "rule",
    "expected_settlement_date",
    "actual_settlement_date",
    "settlement_id",
    "days_overdue",
    "status",
)


def _money(value: Decimal | None) -> str:
    return "" if value is None else f"{value:.2f}"


def _iso(value: object) -> str:
    return "" if value is None else str(value)


def overdue_csv(rows: list[PaymentRow]) -> str:
    """The overdue report: what should have arrived and has not."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(OVERDUE_COLUMNS)
    for row in overdue_rows(rows):
        writer.writerow(
            [
                row.payment_id,
                row.provider,
                _iso(row.payment_date),
                _money(row.amount),
                row.currency,
                _iso(row.expected_settlement_date),
                row.days_overdue,
                row.status.value if row.status else "",
            ]
        )
    return buffer.getvalue()


def all_rows_xlsx(rows: list[PaymentRow]) -> bytes:
    """Every reconciled row as a formatted worksheet."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Settlements"
    sheet.append(list(ALL_COLUMNS))

    header_fill = PatternFill("solid", fgColor="1F2937")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")

    overdue_fill = PatternFill("solid", fgColor="FEE2E2")
    late_fill = PatternFill("solid", fgColor="FEF3C7")

    for row in rows:
        sheet.append(
            [
                row.payment_id,
                row.provider,
                _iso(row.payment_date),
                row.amount,  # openpyxl writes Decimal as an exact numeric cell
                row.currency,
                row.rule_label,
                _iso(row.expected_settlement_date),
                _iso(row.actual_settlement_date),
                row.settlement_id or "",
                row.days_overdue,
                row.status.value if row.status else "NO_RULE",
            ]
        )
        status = row.status.value if row.status else ""
        if status == "OVERDUE":
            for cell in sheet[sheet.max_row]:
                cell.fill = overdue_fill
        elif status == "SETTLED_LATE":
            for cell in sheet[sheet.max_row]:
                cell.fill = late_fill

    for column, width in zip(
        "ABCDEFGHIJK", (16, 12, 14, 12, 10, 22, 22, 20, 18, 13, 14), strict=False
    ):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
