"""Overdue CSV and XLSX exports."""

from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.services.export_service import OVERDUE_COLUMNS, all_rows_xlsx, overdue_csv
from app.services.reconciliation import build_rows
from tests.conftest import AS_OF
from tests.test_reconciliation import add_payment, add_settlement


def test_overdue_csv_contains_only_overdue_rows(demo_rules: Session) -> None:
    session = demo_rules
    add_payment(session, "p_overdue", "PSP_B", date(2026, 9, 14), "30.00", "EUR")  # -> 15 Sep
    add_payment(session, "p_due_today", "PSP_B", date(2026, 9, 16), "20.00", "EUR")
    add_payment(session, "p_settled", "PSP_B", date(2026, 9, 10), "40.00", "EUR")
    add_settlement(session, "s1", "p_settled", date(2026, 9, 11), "40.00", "EUR")

    content = overdue_csv(build_rows(session, AS_OF))
    rows = list(csv.DictReader(io.StringIO(content)))

    assert list(rows[0].keys()) == list(OVERDUE_COLUMNS)
    assert len(rows) == 1
    assert rows[0]["payment_id"] == "p_overdue"
    assert rows[0]["provider"] == "PSP_B"
    assert rows[0]["amount"] == "30.00"
    assert rows[0]["currency"] == "EUR"
    assert rows[0]["expected_settlement_date"] == "2026-09-15"
    assert rows[0]["days_overdue"] == "2"
    assert rows[0]["status"] == "OVERDUE"


def test_overdue_csv_has_a_header_even_when_empty(demo_rules: Session) -> None:
    content = overdue_csv(build_rows(demo_rules, AS_OF))
    assert content.strip() == ",".join(OVERDUE_COLUMNS)


def test_xlsx_export_keeps_exact_amounts(demo_rules: Session) -> None:
    session = demo_rules
    add_payment(session, "p1", "PSP_B", date(2026, 9, 14), "1234.56", "EUR")
    payload = all_rows_xlsx(build_rows(session, AS_OF))
    sheet = load_workbook(io.BytesIO(payload)).active
    header = [cell.value for cell in sheet[1]]
    assert header[0] == "payment_id"
    assert sheet.cell(row=2, column=1).value == "p1"
    assert Decimal(str(sheet.cell(row=2, column=4).value)) == Decimal("1234.56")
    assert sheet.cell(row=2, column=11).value == "OVERDUE"
